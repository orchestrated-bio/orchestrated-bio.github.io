#!/usr/bin/env python3
"""Create website previews from retained public-data selections, read-only."""
import argparse
import csv
import gzip
import hashlib
import html
import json
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import openpyxl
from PIL import ImageFont


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def wrap(value, width, size):
    font = ImageFont.truetype("/System/Library/Fonts/Supplemental/Arial.ttf", size)
    words = str(value if value is not None else "").split()
    lines, line = [], ""
    for word in words:
        candidate = (line + " " + word).strip()
        if line and font.getlength(candidate) > width:
            lines.append(line)
            line = word
        else:
            line = candidate
    if line:
        lines.append(line)
    return lines


def sheet_preview(sheet, first, last, widths, output, *, row_indices=None, columns=None):
    row_indices = row_indices or list(range(first, last + 1))
    columns = columns or list(range(1, len(widths) + 1))
    rows = [[sheet.cell(r, c).value for c in columns] for r in row_indices]
    gutter, start_y, header_h = 40, 88, 90
    row_h = 90 if sheet.title == "Sources" else 120
    width = gutter + sum(widths)
    height = start_y + header_h + (len(rows) - 1) * row_h + 60
    parts = [f'<svg xmlns="http://www.w3.org/2000/svg" width="{width}" height="{height}" viewBox="0 0 {width} {height}">',
             f'<title>Evidence.xlsx: {sheet.title}, rows {first} to {last}</title>',
             f'<rect width="{width}" height="{height}" fill="#fff"/>',
             f'<rect width="{width}" height="48" fill="#214f41"/>',
             f'<text x="22" y="31" fill="#fff" font-family="Arial,sans-serif" font-size="21">Evidence.xlsx</text>',
             f'<rect y="48" width="{width}" height="40" fill="#f2f5f3"/>']
    x = gutter
    for i, col_w in enumerate(widths):
        letter = openpyxl.utils.get_column_letter(columns[i])
        parts.append(f'<text x="{x + col_w / 2}" y="74" text-anchor="middle" font-family="Arial,sans-serif" font-size="17" fill="#55675f">{letter}</text>')
        x += col_w
    for r, row in enumerate(rows):
        y, h = start_y + (0 if r == 0 else header_h + (r-1) * row_h), header_h if r == 0 else row_h
        fill = "#e7f0eb" if r == 0 else ("#f4f8f5" if r % 2 == 0 else "#fff")
        parts.append(f'<rect x="{gutter}" y="{y}" width="{sum(widths)}" height="{h}" fill="{fill}"/>')
        parts.append(f'<text x="20" y="{y+h/2+6}" text-anchor="middle" fill="#718078" font-family="Arial,sans-serif" font-size="16">{row_indices[r]}</text>')
        x = gutter
        for c, value in enumerate(row):
            size = 18 if r == 0 else 20
            is_number = isinstance(value, (int, float))
            text = f"{value:g}" if is_number else str(value if value is not None else "")
            lines = wrap(text, widths[c]-24, size)
            line_h = size * 1.38
            text_y = y + (h - len(lines)*line_h)/2 + size
            for line in lines:
                align = "end" if is_number and r else "start"
                text_x = x + widths[c] - 12 if align == "end" else x + 12
                weight = ' font-weight="600"' if r == 0 else ""
                parts.append(f'<text x="{text_x}" y="{text_y:.1f}" text-anchor="{align}" font-family="Arial,sans-serif" font-size="{size}" fill="#213e32"{weight}>{html.escape(line)}</text>')
                text_y += line_h
            parts.append(f'<path d="M{x} {y}v{h}" stroke="#dce5df"/>')
            x += widths[c]
        parts.append(f'<path d="M{gutter} {y+h}H{width}" stroke="#dce5df"/>')
    selection = (f"{sheet.title}!A{first}:{chr(64+len(widths))}{last}" if columns == list(range(1, len(widths) + 1))
                 else f"{sheet.title}: rows {', '.join(map(str, row_indices[1:]))}; columns {', '.join(openpyxl.utils.get_column_letter(c) for c in columns)}")
    parts.append(f'<text x="{gutter+12}" y="{height-20}" font-family="Arial,sans-serif" font-size="17" fill="#406552">{selection}</text></svg>')
    output.write_text("\n".join(parts))
    return {"sheet": sheet.title, "range": selection.split("!", 1)[-1], "rows": row_indices, "columns": columns, "values": rows}


def umap_preview(metadata, output):
    with gzip.open(metadata, "rt") as handle:
        records = list(csv.DictReader(handle))
    xy = np.array([[float(row["UMAP_1"]), float(row["UMAP_2"])] for row in records])
    labels = np.array([row["celltype"] for row in records])
    palette = ["#5a9f8c", "#7da7ce", "#dba953", "#899ec1", "#c88799", "#9fab76", "#69afbd", "#bc956e", "#84987e", "#b4a0bc", "#6091a1", "#bd8192", "#bdac75", "#8aacc1", "#91b9a3"]
    fig, ax = plt.subplots(figsize=(10, 6.1), dpi=210)
    fig.patch.set_facecolor("#fbfcfb")
    ax.set_facecolor("#fbfcfb")
    for label, color in zip(sorted(set(labels)), palette):
        points = xy[labels == label]
        ax.scatter(points[:, 0], points[:, 1], s=.55, alpha=.6, color=color, linewidths=0, rasterized=True)
        center = np.median(points, axis=0)
        if label == "URO1":
            center += [-.9, .6]
        if label == "URO2":
            center += [.8, .5]
        ax.text(*center, label, fontsize=9, ha="center", va="center", weight="bold", color="#203c31",
                bbox={"facecolor": "#fbfcfb", "alpha": .82, "edgecolor": "none", "pad": 2})
    ax.set(xlabel="UMAP 1", ylabel="UMAP 2")
    ax.set_xticks([])
    ax.set_yticks([])
    ax.spines[["top", "right"]].set_visible(False)
    ax.spines[["bottom", "left"]].set_color("#ced8d1")
    ax.xaxis.label.set_color("#66786e")
    ax.yaxis.label.set_color("#66786e")
    fig.suptitle("Human kidney cell populations", x=.095, y=.965, ha="left", fontsize=19, weight="bold", color="#213e32")
    fig.text(.095, .90, "Published UMAP · GSE185948 · 102,710 nuclei", fontsize=11, color="#5d7165")
    fig.subplots_adjust(left=.095, right=.965, top=.85, bottom=.11)
    fig.savefig(output, facecolor=fig.get_facecolor())
    plt.close(fig)
    return {"n_nuclei": len(records), "coordinates": "Published UMAP_1 and UMAP_2", "labels": sorted(set(labels)), "purpose": "Descriptive kidney cell map from the single-nucleus follow-on"}


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--workbook", type=Path, required=True)
    parser.add_argument("--metadata", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    args.output.mkdir(parents=True, exist_ok=True)
    original_hash = sha(args.workbook)
    workbook = openpyxl.load_workbook(args.workbook, read_only=True, data_only=True)
    selections = []
    for name, lo, hi, widths in [("Summary", 5, 8, [170, 245, 80, 205, 115, 115]),
                                 ("Parameters", 1, 4, [235, 140, 265, 90, 205]),
                                 ("Sources", 1, 5, [180, 125, 170, 80, 380])]:
        selections.append(sheet_preview(workbook[name], lo, hi, widths, args.output / f"workbook-{name.lower()}.svg"))
    selections.append(sheet_preview(workbook["Parameters"], 1, 28, [150, 170, 190, 190, 190], args.output / "workbook-expression.svg",
                                    row_indices=[1, 26, 25, 28], columns=[3, 4, 12, 13, 20]))
    workbook.close()
    assert sha(args.workbook) == original_hash
    umap = umap_preview(args.metadata, args.output / "single-nucleus-umap.png")
    provenance = {"workbook_sha256": original_hash, "workbook_name": "Evidence.xlsx", "selections": selections,
                  "umap": {"source": "GSE185948_metadata_RNA.csv.gz", "sha256": sha(args.metadata), **umap}}
    (args.output / "preview-sources.json").write_text(json.dumps(provenance, indent=2))
    print("Created four reflowed workbook previews and a plot of the published UMAP coordinates.")


if __name__ == "__main__":
    main()
